"""產生試行工具包的三份 Excel：採購清單、既有案場盤點表、清坑秤砂與試點紀錄表。

用法：cd print && npm run kit（會先匯出 tools/budget.json）後執行
      python3 tools/build_xlsx.py
需要 openpyxl（pip install openpyxl）。
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
KIT = HERE.parent.parent / "試行工具包"
FONT = "Microsoft JhengHei"
GREEN = "0F6E56"
SOFT = "E6F1ED"
thin = Side(style="thin", color="D9D7D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, text, sub, width):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=GREEN)
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, size=10, color="4A4D53")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.row_dimensions[1].height = 26


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, r0, r1, ncol):
    for r in range(r0, r1 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT, size=10)


def purchase_list():
    phases = json.loads((HERE / "budget.json").read_text(encoding="utf-8"))
    wb = Workbook()
    first = True
    for ph in phases:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = ph["name"]
        cols = ["✓ 已買", "品項", "計入預算", "最低 (NT$)", "最高 (NT$)", "實際花費 (NT$)", "價格來源", "參考／通路", "備註"]
        widths = [8, 52, 10, 12, 12, 14, 10, 34, 30]
        title(ws, f"{ph['name']}：{ph['title']}（{ph['when']}）", f"目標：{ph['goal']}。價格為 2026-09 查價；「估計」下單前請再詢價。", len(cols))
        header(ws, 4, cols, widths)
        r = 5
        for it in ph["items"]:
            ws.cell(row=r, column=1, value="")
            ws.cell(row=r, column=2, value=it["name"])
            ws.cell(row=r, column=3, value="否" if it.get("optional") else "是")
            ws.cell(row=r, column=4, value=it["lo"])
            ws.cell(row=r, column=5, value=it["hi"])
            ws.cell(row=r, column=7, value="查證" if it["src"] == "verified" else "估計")
            ws.cell(row=r, column=8, value=it.get("ref", ""))
            ws.cell(row=r, column=9, value=it.get("note", ""))
            for c in (4, 5, 6):
                ws.cell(row=r, column=c).number_format = "#,##0"
            r += 1
        last = r - 1
        body(ws, 5, last, len(cols))
        dv = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"C5:C{last}")
        dv2 = DataValidation(type="list", formula1='"✓,"', allow_blank=True)
        ws.add_data_validation(dv2)
        dv2.add(f"A5:A{last}")
        # 合計（只加「計入預算＝是」）
        ws.cell(row=r + 1, column=2, value="小計（計入預算＝是）").font = Font(name=FONT, bold=True)
        for c in (4, 5):
            L = get_column_letter(c)
            cell = ws.cell(row=r + 1, column=c, value=f'=SUMIF($C$5:$C${last},"是",{L}5:{L}{last})')
            cell.font = Font(name=FONT, bold=True)
            cell.number_format = "#,##0"
            cell.fill = PatternFill("solid", fgColor=SOFT)
        cell = ws.cell(row=r + 1, column=6, value=f"=SUM(F5:F{last})")
        cell.font = Font(name=FONT, bold=True)
        cell.number_format = "#,##0"
        cell.fill = PatternFill("solid", fgColor=SOFT)
    wb.save(KIT / "01-第0階段採購清單.xlsx")


def site_inventory():
    wb = Workbook()
    ws = wb.active
    ws.title = "既有案場"
    cols = ["#", "案名", "完工年份", "地址／區域", "類型", "物業公司／管理單位", "聯絡人", "電話／LINE", "入口寬 (m)", "入口深 (m)", "室內／室外", "面材", "上次換毯", "與我們關係", "日流量估計", "試點適合度 (1–5)", "下一步", "備註"]
    widths = [5, 24, 9, 20, 11, 20, 12, 16, 9, 9, 9, 14, 11, 11, 11, 11, 18, 24]
    title(ws, "既有案場盤點表", "從官網案例與出貨紀錄整理。適合度：關係好、有物業公司駐點、日流量大、室外或雨天多 → 分數高。先挑 5–8 個 4 分以上的案場談試點。", len(cols))
    header(ws, 4, cols, widths)
    examples = [
        [1, "（範例）○○社區大廳", 2019, "台北市信義區", "住宅社區", "○○物業", "王總幹事", "", 2.4, 1.5, "室內", "碳黑浪花紋", 2019, "良好", 800, 4, "寄免費健檢", "這一列是範例，可以刪除"],
    ]
    for i, row in enumerate(examples):
        for j, v in enumerate(row, start=1):
            ws.cell(row=5 + i, column=j, value=v)
    for r in range(6, 105):
        ws.cell(row=r, column=1, value=r - 4)
    body(ws, 5, 104, len(cols))
    for name, col, opts in [
        ("類型", 5, "住宅社區,商辦,商場,醫院,科技廠,飯店,學校場館,公部門,其他"),
        ("室內外", 11, "室內,室外,風除室"),
        ("關係", 14, "良好,普通,很久沒聯絡"),
        ("下一步", 17, "寄免費健檢,電話聯絡,已拜訪,談試點,已簽試點,暫不適合"),
    ]:
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        L = get_column_letter(col)
        dv.add(f"{L}5:{L}104")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("P5:P104")
    ws.conditional_formatting.add("P5:P104", CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=PatternFill("solid", fgColor="CDEBDD")))
    ws2 = wb.create_sheet("評分方式")
    rows = [
        ("項目", "加分條件"),
        ("關係", "跟業主、物業或建商窗口還有聯絡"),
        ("物業公司", "有大型物業公司駐點（可以一起談分潤）"),
        ("流量", "日流量大（1,000 人次以上）"),
        ("環境", "室外或風除室、雨天積水多、曾經有人滑倒"),
        ("案例價值", "醫院、機場、科技廠、知名建商：對其他客戶有說服力"),
        ("施工", "地墊狀況良好，只需換底座條就能改裝"),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ws2.cell(row=i, column=1, value=a).font = Font(name=FONT, bold=i == 1)
        ws2.cell(row=i, column=2, value=b).font = Font(name=FONT, bold=i == 1)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 60
    wb.save(KIT / "02-既有案場盤點表.xlsx")


def pilot_log():
    wb = Workbook()
    ws = wb.active
    ws.title = "清坑秤砂"
    cols = ["日期", "案場", "入口", "距上次清坑（天）", "期間人次（系統）", "期間是否下雨", "秤到的砂重 (g)", "每千人次砂重 (g)", "模型估計 (g)", "誤差 %", "處理人", "備註"]
    widths = [12, 20, 12, 12, 14, 12, 14, 14, 13, 10, 12, 24]
    title(ws, "清坑秤砂紀錄", "每次清坑都秤一次砂（吊秤即可）。這份數據用來校正積砂模型，是別人沒有的資產。", len(cols))
    header(ws, 4, cols, widths)
    for r in range(5, 205):
        ws.cell(row=r, column=8, value=f'=IF(AND(E{r}>0,G{r}<>""),G{r}/E{r}*1000,"")').number_format = "#,##0.0"
        ws.cell(row=r, column=10, value=f'=IF(AND(I{r}>0,G{r}<>""),(I{r}-G{r})/G{r},"")').number_format = "0%"
        ws.cell(row=r, column=1).number_format = "yyyy/mm/dd"
    body(ws, 5, 204, len(cols))
    dv = DataValidation(type="list", formula1='"晴,雨,晴雨都有"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F5:F204")

    ws2 = wb.create_sheet("試點事件紀錄")
    cols2 = ["日期時間", "案場", "入口", "事件", "系統有沒有通知", "多久有人處理（分）", "處理人", "結果", "備註"]
    widths2 = [16, 20, 12, 16, 12, 14, 12, 20, 30]
    title(ws2, "試點事件紀錄", "記下每一次警報、滑倒、客訴、設備異常。試點結束時用來算成功指標。", len(cols2))
    header(ws2, 4, cols2, widths2)
    body(ws2, 5, 304, len(cols2))
    dv2 = DataValidation(type="list", formula1='"積砂清潔,濕滑警示,積水,夜間入侵,毯條更換,設備異常,客訴,滑倒事件,其他"', allow_blank=True)
    ws2.add_data_validation(dv2)
    dv2.add("D5:D304")
    dv3 = DataValidation(type="list", formula1='"有,沒有,不適用"', allow_blank=True)
    ws2.add_data_validation(dv3)
    dv3.add("E5:E304")

    ws3 = wb.create_sheet("計步驗證")
    cols3 = ["日期", "時段", "影片人工計數（人）", "系統計數（人）", "誤差 %", "方向正確數", "方向錯誤數", "方向正確率", "備註（並排、推車等）"]
    widths3 = [12, 14, 16, 14, 10, 12, 12, 12, 30]
    title(ws3, "計步驗證（第 0 階段）", "公司門口放一週，錄影人工計數當真值。過關標準：誤差 < 5%、方向正確率 > 95%。", len(cols3))
    header(ws3, 4, cols3, widths3)
    for r in range(5, 105):
        ws3.cell(row=r, column=5, value=f'=IF(C{r}>0,(D{r}-C{r})/C{r},"")').number_format = "0.0%"
        ws3.cell(row=r, column=8, value=f'=IF((F{r}+G{r})>0,F{r}/(F{r}+G{r}),"")').number_format = "0.0%"
    body(ws3, 5, 104, len(cols3))
    wb.save(KIT / "06-清坑秤砂與試點紀錄表.xlsx")


if __name__ == "__main__":
    KIT.mkdir(exist_ok=True)
    purchase_list()
    site_inventory()
    pilot_log()
    print("✓ xlsx x3 ->", KIT)
